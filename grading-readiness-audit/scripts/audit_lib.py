"""Shared helpers for the gameplay-log -> progress-point grading audit.

Every phase script imports from here: config/paths, gameplay-log loading
(with client-timestamp ordering), event-key parsing, docx table extraction,
dialogue-reference loading, and output writers.
"""

import csv
import json
import os
import re
import shutil
import sys
import tempfile
import zipfile
from xml.etree import ElementTree

# ---------------------------------------------------------------------------
# Paths / config
# ---------------------------------------------------------------------------

AUDIT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPO_ROOT = os.path.dirname(AUDIT_DIR)


def utf8_stdout():
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass


def load_config(path=None):
    import yaml

    if path is None:
        path = os.environ.get("PPAUDIT_CONFIG") or os.path.join(
            AUDIT_DIR, "config", "audit-config.yaml"
        )
    with open(path, encoding="utf-8") as f:
        cfg = yaml.safe_load(f)
    return cfg


def repo_path(rel):
    return rel if os.path.isabs(rel) else os.path.join(REPO_ROOT, rel)


def out_path(cfg, *parts):
    p = os.path.join(AUDIT_DIR, cfg.get("outputs_dir", "outputs"), *parts)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    return p


def report_path(cfg, *parts):
    p = os.path.join(AUDIT_DIR, cfg.get("reports_dir", "reports"), *parts)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    return p


def write_json(path, obj):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, indent=2, ensure_ascii=False)


def read_json(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def write_csv(path, rows, fieldnames):
    # utf-8-sig so Excel opens the file cleanly.
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)


# ---------------------------------------------------------------------------
# Gameplay-log loading
# ---------------------------------------------------------------------------


def _oid(doc):
    v = doc.get("_id")
    if isinstance(v, dict):
        return v.get("$oid") or v.get("$id") or ""
    return v if v is not None else ""


def load_logs(cfg):
    """Load every *.json array under cfg['log_dir'].

    Returns (records, meta). Records are dicts with `_id` normalized to the
    24-char hex string and `_src` set to the source filename. The list is
    sorted by client `timestamp` (tiebreak `_id`) — gameplay chronology, NOT
    arrival order (the export array is newest-first and `_id` order disagrees
    with client time for a small share of adjacent pairs).
    """
    log_dir = repo_path(cfg["log_dir"])
    files = sorted(
        f for f in os.listdir(log_dir) if f.lower().endswith(".json")
    )
    if not files:
        raise SystemExit(f"No .json log files found in {log_dir}")
    records = []
    parse_errors = []
    for fn in files:
        fp = os.path.join(log_dir, fn)
        try:
            with open(fp, encoding="utf-8") as f:
                arr = json.load(f)
        except Exception as e:  # malformed file: warn, keep going
            parse_errors.append({"file": fn, "error": str(e)})
            continue
        if not isinstance(arr, list):
            parse_errors.append({"file": fn, "error": "top level is not a JSON array"})
            continue
        for d in arr:
            d = dict(d)
            d["_id"] = _oid(d)
            d["_src"] = fn
            records.append(d)
    records.sort(key=lambda d: (d.get("timestamp") or "", d["_id"]))
    # player-field detection
    field = cfg.get("player_field", "auto")
    detected = None
    if field == "auto":
        sample = records[0] if records else {}
        if "playerId" in sample:
            detected = "playerId"
        elif "user_id" in sample:
            detected = "user_id"
        else:
            detected = "user_id"
    else:
        detected = field
    players = sorted({r.get(detected) for r in records if r.get(detected)})
    meta = {
        "log_dir": cfg["log_dir"],
        "files": files,
        "parse_errors": parse_errors,
        "record_count": len(records),
        "player_field": detected,
        "player_field_mode": field,
        "players": players,
        "versions": sorted({r.get("version") for r in records if r.get("version")}),
    }
    return records, meta


# ---------------------------------------------------------------------------
# Event-key helpers
# ---------------------------------------------------------------------------

DIALOGUE_KEY_RE = re.compile(r"^DialogueNodeEvent:(\d+):(\d+)$")
QUEST_KEY_RE = re.compile(r"^(questActiveEvent|questFinishEvent):(\d+)$")


def parse_event_key(key):
    """Classify an eventKey string.

    Returns ("dialogue", conv, node) / ("quest", kind, qid) / ("other", key, None).
    """
    m = DIALOGUE_KEY_RE.match(key or "")
    if m:
        return ("dialogue", int(m.group(1)), int(m.group(2)))
    m = QUEST_KEY_RE.match(key or "")
    if m:
        return ("quest", m.group(1), int(m.group(2)))
    return ("other", key, None)


# ---------------------------------------------------------------------------
# Docx table extraction
# ---------------------------------------------------------------------------

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _readable_copy(path):
    """OneDrive files-on-demand placeholders raise PermissionError on open;
    copying hydrates them. Python itself cannot trigger hydration in this
    environment, but PowerShell's Copy-Item can, so fall back to that.
    Returns a path that is safe to open."""
    try:
        with open(path, "rb"):
            return path, None
    except PermissionError:
        tmp = os.path.join(tempfile.gettempdir(), "ppaudit-" + os.path.basename(path))
        try:
            shutil.copy(path, tmp)
        except PermissionError:
            import subprocess

            subprocess.run(
                [
                    "powershell.exe",
                    "-NoProfile",
                    "-Command",
                    f'Copy-Item -LiteralPath "{path}" -Destination "{tmp}" -Force',
                ],
                check=True,
                capture_output=True,
            )
        return tmp, tmp


def docx_tables(path):
    """Yield tables from a .docx as lists of rows, each row a list of cell
    texts (paragraphs inside a cell joined by newlines)."""
    readable, _tmp = _readable_copy(path)
    with zipfile.ZipFile(readable) as z:
        root = ElementTree.fromstring(z.read("word/document.xml"))
    tables = []
    for tbl in root.iter(_W + "tbl"):
        rows = []
        for tr in tbl.findall(_W + "tr"):
            cells = []
            for tc in tr.findall(_W + "tc"):
                paras = []
                for p in tc.iter(_W + "p"):
                    text = "".join(t.text or "" for t in p.iter(_W + "t"))
                    paras.append(text)
                cells.append("\n".join(paras).strip())
            rows.append(cells)
        tables.append(rows)
    return tables


def docx_paragraphs(path):
    """Top-level paragraph texts (used to find the 'Unit N' headings that
    precede each table in Progress-Points.docx)."""
    readable, _tmp = _readable_copy(path)
    with zipfile.ZipFile(readable) as z:
        root = ElementTree.fromstring(z.read("word/document.xml"))
    body = root.find(_W + "body")
    items = []  # ("p", text) | ("tbl", index)
    tbl_i = 0
    for child in body:
        if child.tag == _W + "p":
            text = "".join(t.text or "" for t in child.iter(_W + "t")).strip()
            items.append(("p", text))
        elif child.tag == _W + "tbl":
            items.append(("tbl", tbl_i))
            tbl_i += 1
    return items


# ---------------------------------------------------------------------------
# Dialogue references
# ---------------------------------------------------------------------------


_EXPORT_DATE_RE = re.compile(r"\d{4}-\d{2}-\d{2}")


def dialogue_export_label(cfg):
    """Label for the configured Unity dialogue export: the YYYY-MM-DD in its
    file name (e.g. '2026-09-21'), else the bare file name. Report and
    reconciliation wording is built from this so switching
    `dialogue_export_csv` in audit-config.yaml never leaves a stale date."""
    name = os.path.basename(cfg.get("dialogue_export_csv") or "")
    m = _EXPORT_DATE_RE.search(name)
    return m.group(0) if m else (name or "unknown")


def load_dialogue_xlsx(cfg):
    """(conv, node) -> text map from the workbook named by `dialogue_xlsx`
    (Dialogue-ID-Texts.xlsx; replaced 2026-09-21 to match build 20260914-,
    the previous copy is kept as Dialogue-ID-Texts-Old.xlsx)."""
    import openpyxl

    path, _ = _readable_copy(repo_path(cfg["dialogue_xlsx"]))
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    out = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        conv, node, text = (row + (None,) * 3)[:3]
        if conv is None or node is None:
            continue
        try:
            out[(int(conv), int(node))] = (text or "").strip()
        except (TypeError, ValueError):
            continue
    return out


def load_dialogue_export(cfg):
    """Unity dialogue-database export named by `dialogue_export_csv`
    (2026-09-21-MHSDialogueExport.csv since 2026-09-23; sections Database /
    Conversations / DialogueEntries / OutgoingLinks).

    Returns (conversations, entries):
      conversations: conv_id -> title
      entries: (conv, node) -> {"title","menu_text","dialogue_text","entrytag"}
    """
    path, _ = _readable_copy(repo_path(cfg["dialogue_export_csv"]))
    with open(path, encoding="utf-8-sig", errors="replace", newline="") as f:
        rows = list(csv.reader(f))
    sections = {}
    for i, row in enumerate(rows):
        if len(row) == 1 and row[0] and not row[0].isdigit():
            sections[row[0]] = i
    conversations = {}
    entries = {}
    conv_start = sections.get("Conversations")
    if conv_start is not None:
        for row in rows[conv_start + 3:]:
            if len(row) <= 1:
                break
            if row[0].isdigit():
                conversations[int(row[0])] = row[1]
    de_start = sections.get("DialogueEntries")
    if de_start is not None:
        header = rows[de_start + 1]
        idx = {name: i for i, name in enumerate(header)}
        for row in rows[de_start + 3:]:
            if len(row) <= 1 or (len(row) == 1 and not row[0].isdigit()):
                break
            if len(row) < len(header):
                continue
            try:
                conv = int(row[idx["ConvID"]])
                node = int(row[idx["ID"]])
            except (ValueError, KeyError):
                continue
            entries[(conv, node)] = {
                "entrytag": row[idx.get("entrytag", 0)],
                "title": row[idx["Title"]],
                "menu_text": row[idx["MenuText"]],
                "dialogue_text": row[idx["DialogueText"]],
            }
    return conversations, entries


_WS_RE = re.compile(r"\s+")
_TAG_RE = re.compile(r"\[[^\]]*\]")  # [var=...], [press M/...] style tokens


def norm_text(s):
    """Normalization used for dialogue-text comparison: strip markup tokens,
    collapse whitespace, casefold, unify curly quotes."""
    if not s:
        return ""
    s = s.replace("’", "'").replace("‘", "'")
    s = s.replace("“", '"').replace("”", '"')
    s = s.replace("…", "...")
    s = _TAG_RE.sub(" ", s)
    return _WS_RE.sub(" ", s).strip().casefold()
